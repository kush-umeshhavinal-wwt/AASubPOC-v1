"""
SharePoint/Excel Client for Comment Operations
Handles SharePoint API integration and Excel file manipulation
"""

import openpyxl
import logging
from typing import List, Dict
import os
import posixpath
from office365.sharepoint.client_context import ClientContext
from comment_sync import config

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class SharePointClient:
    """Client for interacting with SharePoint and Excel files"""
    
    def __init__(
        self,
        site_url: str = None,
        username: str = None,
        password: str = None,
        auth_mode: str = None,
        tenant: str = None,
        client_id: str = None,
        file_path: str = None,
        context=None
    ):
        """
        Initialize SharePoint client
        
        Args:
            site_url: SharePoint site URL (uses config if not provided)
            username: SharePoint username (uses config if not provided)
            password: SharePoint password (uses config if not provided)
        """
        self.site_url = site_url or config.SHAREPOINT_SITE_URL
        self.username = username or config.SHAREPOINT_USERNAME
        self.password = password or config.SHAREPOINT_PASSWORD
        self.auth_mode = (auth_mode or config.SHAREPOINT_AUTH_MODE).lower()
        self.tenant = tenant or config.SHAREPOINT_TENANT
        self.client_id = client_id or config.SHAREPOINT_CLIENT_ID
        self.file_path = file_path or config.SHAREPOINT_FILE_PATH
        self._context = context

        if self.auth_mode == "local":
            self.sharepoint_enabled = False
        elif self.auth_mode == "interactive":
            required = [self.site_url, self.tenant, self.client_id, self.file_path]
            if not all(required):
                raise ValueError("Interactive SharePoint authentication requires site URL, tenant, client ID, and file path")
            self.sharepoint_enabled = True
        elif self.auth_mode == "user_credentials":
            required = [self.site_url, self.username, self.password, self.file_path]
            if not all(required):
                raise ValueError("SharePoint user credentials authentication requires site URL, username, password, and file path")
            self.sharepoint_enabled = True
        else:
            raise ValueError(f"Unsupported SharePoint authentication mode: {self.auth_mode}")

    def _get_context(self):
        if self._context is None:
            context = ClientContext(self.site_url)
            if self.auth_mode == "interactive":
                self._context = context.with_interactive(self.tenant, self.client_id)
            else:
                self._context = context.with_user_credentials(self.username, self.password)
        return self._context
    
    def connect_to_sharepoint(self) -> bool:
        """
        Test connection to SharePoint
        
        Returns:
            bool: True if connection successful
        """
        if not self.sharepoint_enabled:
            logger.info("SharePoint integration not configured")
            return False
        
        try:
            # Test connection by making a simple request
            self._get_context().web.get().execute_query()
            logger.info(f"Connected to SharePoint at: {self.site_url}")
            return True
        except Exception as e:
            logger.error(f"Failed to connect to SharePoint: {e}")
            return False
    
    def download_excel_file(self, file_path: str = None, local_path: str = None) -> str:
        """
        Download Excel file from SharePoint
        
        Args:
            file_path: SharePoint server-relative file path (uses config if not provided)
            local_path: Local path to save the file (uses config if not provided)
            
        Returns:
            Local file path
        """
        if not self.sharepoint_enabled:
            logger.warning("SharePoint integration not enabled - using local file")
            return local_path or config.LOCAL_EXCEL_FILE
        
        file_path = file_path or self.file_path
        local_path = local_path or config.LOCAL_EXCEL_FILE
        
        if not file_path:
            raise ValueError("SharePoint file path is required")
        
        temp_path = f"{local_path}.download"
        try:
            # Ensure directory exists
            local_dir = os.path.dirname(local_path)
            if local_dir:
                os.makedirs(local_dir, exist_ok=True)

            with open(temp_path, 'wb') as local_file:
                remote_file = self._get_context().web.get_file_by_server_relative_path(file_path)
                remote_file.download(local_file).execute_query()
            os.replace(temp_path, local_path)

            logger.info(f"Downloaded Excel file to: {local_path}")
            return local_path
        except Exception as e:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            logger.error(f"Failed to download Excel file from SharePoint: {e}")
            raise
    
    def upload_excel_file(self, file_path: str = None, local_path: str = None) -> bool:
        """
        Upload Excel file to SharePoint
        
        Args:
            file_path: SharePoint server-relative file path (uses config if not provided)
            local_path: Local file path to upload (uses config if not provided)
            
        Returns:
            bool: True if upload successful
        """
        if not self.sharepoint_enabled:
            logger.warning("SharePoint integration not enabled - skipping upload")
            return False
        
        file_path = file_path or self.file_path
        local_path = local_path or config.LOCAL_EXCEL_FILE
        
        if not file_path or not local_path:
            raise ValueError("Both SharePoint file path and local path are required")
        
        try:
            folder_path, file_name = posixpath.split(file_path)
            with open(local_path, 'rb') as local_file:
                content = local_file.read()
            folder = self._get_context().web.get_folder_by_server_relative_path(folder_path)
            folder.upload_file(file_name, content).execute_query()

            logger.info(f"Uploaded Excel file to SharePoint: {file_path}")
            return True
        except Exception as e:
            logger.error(f"Failed to upload Excel file to SharePoint: {e}")
            raise
    
    def read_excel_comments(self, local_path: str = None) -> List[Dict]:
        """
        Read comments from Excel file
        
        Args:
            local_path: Local Excel file path (uses config if not provided)
            
        Returns:
            List of comment dictionaries
        """
        local_path = local_path or config.LOCAL_EXCEL_FILE
        
        if not os.path.exists(local_path):
            raise FileNotFoundError(f"Excel file not found: {local_path}")
        
        try:
            with open(local_path, 'rb') as workbook_file:
                workbook = openpyxl.load_workbook(workbook_file)
            
            # Get the Comments sheet
            if "Details AUG-26" in workbook.sheetnames and config.EXCEL_SHEET_NAME not in workbook.sheetnames:
                raise ValueError("Use scripts.import_real_excel for the real Details AUG-26 workbook")
            if config.EXCEL_SHEET_NAME in workbook.sheetnames:
                sheet = workbook[config.EXCEL_SHEET_NAME]
            else:
                # Use first sheet if Comments sheet doesn't exist
                sheet = workbook.active
                logger.warning(f"Sheet '{config.EXCEL_SHEET_NAME}' not found, using active sheet")
            
            comments = []
            headers = None
            
            for row_num, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_num == 1:
                    # First row is headers
                    headers = [str(cell) if cell is not None else '' for cell in row]
                    continue
                
                # Skip empty rows
                if not any(row):
                    continue
                
                # Create comment dictionary
                comment = {
                    header: str(value) if value is not None else ''
                    for header, value in zip(headers, row)
                    if header in config.EXCEL_HEADERS
                }
                
                # Only include rows with comment_id
                if comment.get('comment_id'):
                    comments.append(comment)
            
            workbook.close()
            logger.info(f"Read {len(comments)} comments from Excel file")
            return comments
            
        except Exception as e:
            logger.error(f"Failed to read comments from Excel file: {e}")
            raise
    
    def write_excel_comments(self, comments: List[Dict], local_path: str = None) -> str:
        """
        Write comments to Excel file
        
        Args:
            comments: List of comment dictionaries
            local_path: Local Excel file path (uses config if not provided)
            
        Returns:
            Local file path
        """
        local_path = local_path or config.LOCAL_EXCEL_FILE
        
        try:
            # Create new workbook
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = config.EXCEL_SHEET_NAME
            
            # Write headers
            sheet.append(config.EXCEL_HEADERS)
            
            # Write comments
            for comment in comments:
                row = [
                    comment.get(header, '')
                    for header in config.EXCEL_HEADERS
                ]
                sheet.append(row)
            
            # Ensure directory exists
            os.makedirs(os.path.dirname(local_path), exist_ok=True)
            
            # Save workbook
            workbook.save(local_path)
            workbook.close()
            
            logger.info(f"Wrote {len(comments)} comments to Excel file: {local_path}")
            return local_path
            
        except Exception as e:
            logger.error(f"Failed to write comments to Excel file: {e}")
            raise


# Convenience functions for easier usage
def connect_to_sharepoint(site_url: str = None, username: str = None, password: str = None) -> SharePointClient:
    """
    Create and return a connected SharePoint client
    
    Args:
        site_url: SharePoint site URL
        username: SharePoint username
        password: SharePoint password
        
    Returns:
        Connected SharePointClient instance
    """
    client = SharePointClient(site_url, username, password)
    if not client.sharepoint_enabled:
        return client
    if not client.connect_to_sharepoint():
        raise ConnectionError(f"Unable to connect to SharePoint site: {client.site_url}")
    return client


if __name__ == "__main__":
    # Test the client
    try:
        client = connect_to_sharepoint()
        print("SharePoint client initialized")
        
        # Test reading from local file
        local_excel = os.path.join(config.PROJECT_ROOT, "sample_data", "comments_sample.xlsx")
        if os.path.exists(local_excel):
            comments = client.read_excel_comments(local_excel)
            print(f"Read {len(comments)} comments from sample Excel file")
            for comment in comments[:3]:
                print(f"  - {comment.get('comment_text', '')[:50]}...")
        
    except Exception as e:
        print(f"Error: {e}")