import argparse
import json
import os
import statistics
from collections import Counter
from datetime import date, datetime, timedelta

import openpyxl
import smartsheet

from comment_sync import config


DEFAULT_SHEET_NAME = "Details AUG-26"
HEADER_ROW = 6
DATA_START_ROW = 7
MAPPED_COLUMNS = {
    "age": 1,
    "aging_bucket": 2,
    "comment_text": 4,
    "pl_name": 6,
    "comment_name": 7,
    "comment_id": 8,
    "sub_program": 20,
    "account": 25,
    "start_date": 37
}
EXPECTED_HEADERS = {
    "age": "Age",
    "aging_bucket": "Aging Bucket",
    "comment_text": "Cost Accountant Comments",
    "comment_name": "Je Batch Name",
    "comment_id": "Je Name"
}
OPTIONAL_FIELDS = {"account", "pl_name", "sub_program"}
EXCEL_EPOCH = date(1899, 12, 30)


def expected_bucket(age):
    if age <= 30:
        return "0-30"
    if age <= 60:
        return "31-60"
    if age <= 90:
        return "61-90"
    if age <= 180:
        return "91-180"
    return "180+"


def excel_serial_to_date(value, source_row):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        raise ValueError(f"Invalid start date at source row {source_row}: {value!r}")
    try:
        serial = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid start date at source row {source_row}: {value!r}") from error
    if serial <= 0:
        raise ValueError(f"Invalid start date serial at source row {source_row}: {value!r}")
    return (EXCEL_EPOCH + timedelta(days=int(serial))).isoformat()


def normalize_age(value, source_row):
    if isinstance(value, bool):
        raise ValueError(f"Invalid age at source row {source_row}: {value!r}")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid age at source row {source_row}: {value!r}") from error
    if not number.is_integer():
        raise ValueError(f"Age must be a whole number at source row {source_row}: {value!r}")
    return int(number)


def load_real_records(file_path, sheet_name=DEFAULT_SHEET_NAME, migration_date=None):
    migration_date = migration_date or date.today()
    workbook_file = open(file_path, "rb")
    workbook = openpyxl.load_workbook(workbook_file, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet {sheet_name!r}. Available sheets: {', '.join(workbook.sheetnames)}")
        sheet = workbook[sheet_name]
        header_labels = {}
        for field, column in MAPPED_COLUMNS.items():
            actual = sheet.cell(HEADER_ROW, column).value
            header_labels[field] = {"column": column, "header": actual}
            expected = EXPECTED_HEADERS.get(field)
            if expected is not None and actual != expected:
                raise ValueError(
                    f"Unexpected header for {field} at row {HEADER_ROW}, column {column}: "
                    f"expected {expected!r}, found {actual!r}"
                )

        occurrence_counts = Counter()
        base_ids = set()
        records = []
        inconsistent_rows = []
        extreme_rows = []
        suffix_mappings = []
        audit_date = migration_date.isoformat()

        for source_row, row in enumerate(
            sheet.iter_rows(min_row=DATA_START_ROW, min_col=1, max_col=37, values_only=True),
            DATA_START_ROW
        ):
            raw_id = row[MAPPED_COLUMNS["comment_id"] - 1]
            if raw_id in (None, ""):
                continue
            mapped = {
                field: row[column - 1]
                for field, column in MAPPED_COLUMNS.items()
            }
            missing = [
                field for field, value in mapped.items()
                if value in (None, "") and field not in OPTIONAL_FIELDS
            ]
            if missing:
                raise ValueError(f"Missing {', '.join(missing)} at source row {source_row}")

            base_id = str(mapped["comment_id"]).strip()
            occurrence_counts[base_id] += 1
            occurrence = occurrence_counts[base_id]
            comment_id = base_id if occurrence == 1 else f"{base_id}#{occurrence}"
            if occurrence > 1:
                suffix_mappings.append({
                    "source_row": source_row,
                    "base_id": base_id,
                    "comment_id": comment_id
                })
            base_ids.add(base_id)
            age = normalize_age(mapped["age"], source_row)
            start_date = excel_serial_to_date(mapped["start_date"], source_row)
            source_bucket = str(mapped["aging_bucket"]).strip()
            bucket = expected_bucket(age)
            if source_bucket != bucket:
                inconsistent_rows.append({
                    "source_row": source_row,
                    "age": age,
                    "aging_bucket": source_bucket,
                    "expected_bucket": bucket
                })
            if age > 1000:
                extreme_rows.append({"source_row": source_row, "age": age})

            records.append({
                "comment_id": comment_id,
                "comment_name": str(mapped["comment_name"]).strip(),
                "comment_text": str(mapped["comment_text"]).strip(),
                "age": age,
                "aging_bucket": bucket,
                "start_date": start_date,
                "account": str(mapped["account"] or "").strip(),
                "pl_name": str(mapped["pl_name"] or "").strip(),
                "sub_program": str(mapped["sub_program"] or "").strip(),
                "created_date": audit_date,
                "modified_date": audit_date,
                "paired_comment_id": "",
                "flag_reason": "",
                "flagged_date": ""
            })
    finally:
        workbook.close()
        workbook_file.close()

    if len({record["comment_id"] for record in records}) != len(records):
        raise ValueError("Generated comment IDs are not unique")

    ages = [record["age"] for record in records]
    report = {
        "file_path": os.path.abspath(file_path),
        "sheet_name": sheet_name,
        "header_row": HEADER_ROW,
        "source_rows": len(records),
        "column_headers": header_labels,
        "base_ids": len(base_ids),
        "suffixed_ids": len(suffix_mappings),
        "suffix_mappings": suffix_mappings,
        "inconsistent_buckets": len(inconsistent_rows),
        "inconsistent_rows": inconsistent_rows,
        "extreme_ages": len(extreme_rows),
        "extreme_rows": extreme_rows,
        "average_age": round(statistics.mean(ages), 1) if ages else 0,
        "median_age": statistics.median(ages) if ages else 0,
        "over_500": sum(age > 500 for age in ages),
        "bucket_180_plus": sum(record["aging_bucket"] == "180+" for record in records)
    }
    return records, report


def serialize_sheet(sheet):
    columns = [{
        "id": column.id,
        "title": column.title,
        "index": column.index,
        "type": str(column.type),
        "primary": bool(getattr(column, "primary", False)),
        "options": list(column.options or []) if getattr(column, "options", None) else []
    } for column in sheet.columns]
    column_titles = {column.id: column.title for column in sheet.columns}
    rows = []
    for row in sheet.rows:
        rows.append({
            "id": row.id,
            "row_number": row.row_number,
            "cells": [{
                "column_id": cell.column_id,
                "column_title": column_titles.get(cell.column_id),
                "value": cell.value,
                "display_value": cell.display_value
            } for cell in row.cells]
        })
    return {"name": sheet.name, "id": sheet.id, "columns": columns, "rows": rows}


def write_json(path, payload):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as output:
        json.dump(payload, output, indent=2, default=str)


def get_row_values(sheet):
    column_titles = {column.id: column.title for column in sheet.columns}
    return {
        row.id: {
            column_titles.get(cell.column_id): cell.value
            for cell in row.cells
            if cell.column_id in column_titles
        }
        for row in sheet.rows
    }


def normalize_field_value(field, value):
    if value in (None, ""):
        return ""
    if field == "age":
        return int(float(value))
    return str(value)


def build_smartsheet_row(record, column_map):
    row = smartsheet.models.Row()
    row.to_bottom = True
    row.cells = [
        smartsheet.models.Cell({
            "columnId": column_map[field],
            "value": record[field]
        })
        for field in config.EXCEL_HEADERS
        if record[field] != ""
    ]
    return row


def chunked(values, size=100):
    for index in range(0, len(values), size):
        yield values[index:index + size]


def create_smartsheet_client():
    client = smartsheet.Smartsheet(config.SMARTSHEET_API_TOKEN)
    client.errors_as_exceptions = True
    return client


def stage_migration(file_path, sheet_name=DEFAULT_SHEET_NAME, sheet_id=None, allow_overlap=False):
    records, report = load_real_records(file_path, sheet_name=sheet_name)
    sheet_id = sheet_id or config.SMARTSHEET_SHEET_ID
    if not sheet_id:
        raise ValueError("Smartsheet sheet ID is required")
    client = create_smartsheet_client()
    sheet = client.Sheets.get_sheet(sheet_id)
    column_map = {column.title: column.id for column in sheet.columns}
    missing = [field for field in config.EXCEL_HEADERS if field not in column_map]
    if missing:
        raise ValueError(f"Smartsheet is missing target columns: {', '.join(missing)}")

    target_ids = {record["comment_id"] for record in records}
    current_values = get_row_values(sheet)
    existing_ids = {
        str(values.get("comment_id"))
        for values in current_values.values()
        if values.get("comment_id") not in (None, "")
    }
    overlap = sorted(target_ids & existing_ids)
    if overlap and not allow_overlap:
        raise ValueError(f"Target IDs already exist in Smartsheet; migration may already be staged: {overlap[:5]}")
    if overlap:
        print(f"Reimport mode: {len(overlap)} existing comment IDs will be replaced at finalize")

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_path = os.path.join(config.LOCAL_TEMP_DIR, f"smartsheet-pre-real-migration-{timestamp}.json")
    state_path = os.path.join(config.LOCAL_TEMP_DIR, f"smartsheet-real-migration-state-{timestamp}.json")
    backup = {
        "created_at": datetime.now().isoformat(),
        "source_report": report,
        "sheet": serialize_sheet(sheet)
    }
    write_json(backup_path, backup)

    old_row_ids = [row.id for row in sheet.rows]
    status_column_id = column_map.get("status")
    created_row_ids = []
    try:
        for record_chunk in chunked(records):
            result = client.Sheets.add_rows(
                sheet_id,
                [build_smartsheet_row(record, column_map) for record in record_chunk]
            )
            created_row_ids.extend(row.id for row in result.data)

        verified = client.Sheets.get_sheet(sheet_id)
        verified_values = get_row_values(verified)
        created_values = [verified_values[row_id] for row_id in created_row_ids if row_id in verified_values]
        created_ids = {str(values.get("comment_id")) for values in created_values}
        if len(created_row_ids) != len(records) or created_ids != target_ids:
            raise RuntimeError(
                f"Staged verification failed: created rows={len(created_row_ids)}, "
                f"expected rows={len(records)}, verified IDs={len(created_ids)}"
            )
        created_by_id = {
            str(values.get("comment_id")): values
            for values in created_values
        }
        mismatches = []
        for record in records:
            actual = created_by_id[record["comment_id"]]
            for field, expected in record.items():
                if normalize_field_value(field, actual.get(field)) != normalize_field_value(field, expected):
                    mismatches.append((record["comment_id"], field, actual.get(field), expected))
        if mismatches:
            raise RuntimeError(f"Staged field verification failed: {mismatches[:5]}")
    except Exception:
        for row_ids in chunked(created_row_ids):
            client.Sheets.delete_rows(sheet_id, row_ids)
        raise

    state = {
        "created_at": datetime.now().isoformat(),
        "sheet_id": str(sheet_id),
        "backup_path": backup_path,
        "source_file": os.path.abspath(file_path),
        "source_sheet": sheet_name,
        "old_row_ids": old_row_ids,
        "new_row_ids": created_row_ids,
        "expected_comment_ids": sorted(target_ids),
        "status_column_id": status_column_id,
        "finalized": False,
        "report": report
    }
    write_json(state_path, state)
    return state_path, state


def load_state(state_path):
    with open(state_path, "r", encoding="utf-8") as source:
        return json.load(source)


def finalize_migration(state_path):
    state = load_state(state_path)
    if state.get("finalized"):
        raise ValueError("Migration state is already finalized")
    client = create_smartsheet_client()
    sheet_id = state["sheet_id"]
    sheet = client.Sheets.get_sheet(sheet_id)
    values = get_row_values(sheet)
    staged_ids = {
        str(values[row_id].get("comment_id"))
        for row_id in state["new_row_ids"]
        if row_id in values
    }
    if staged_ids != set(state["expected_comment_ids"]):
        raise RuntimeError("Staged rows no longer match the migration state; refusing to delete old rows")

    for row_ids in chunked(state["old_row_ids"]):
        client.Sheets.delete_rows(sheet_id, row_ids)
    if state.get("status_column_id"):
        client.Sheets.delete_column(sheet_id, state["status_column_id"])

    verified = client.Sheets.get_sheet(sheet_id)
    column_titles = {column.title for column in verified.columns}
    final_values = get_row_values(verified)
    final_ids = {
        str(values.get("comment_id"))
        for values in final_values.values()
        if values.get("comment_id") not in (None, "")
    }
    if final_ids != set(state["expected_comment_ids"]):
        raise RuntimeError("Final Smartsheet IDs do not match the expected real dataset")
    if "status" in column_titles:
        raise RuntimeError("Status column still exists after finalization")
    missing = [field for field in config.EXCEL_HEADERS if field not in column_titles]
    if missing:
        raise RuntimeError(f"Final Smartsheet is missing columns: {missing}")

    state["finalized"] = True
    state["finalized_at"] = datetime.now().isoformat()
    write_json(state_path, state)
    return state


def rollback_staged_migration(state_path):
    state = load_state(state_path)
    if state.get("finalized"):
        raise ValueError("Cannot roll back staged rows after finalization")
    client = create_smartsheet_client()
    for row_ids in chunked(state["new_row_ids"]):
        client.Sheets.delete_rows(state["sheet_id"], row_ids)
    state["rolled_back"] = True
    state["rolled_back_at"] = datetime.now().isoformat()
    write_json(state_path, state)
    return state


def print_report(report):
    print(f"Source: {report['file_path']}")
    print(f"Worksheet: {report['sheet_name']}")
    print(f"Rows to import: {report['source_rows']}")
    print(f"Base Column H IDs: {report['base_ids']}")
    print(f"Suffixed duplicate occurrences: {report['suffixed_ids']}")
    print(f"Average age: {report['average_age']}")
    print(f"Median age: {report['median_age']}")
    print(f"Over 500 days: {report['over_500']}")
    print(f"180+ bucket: {report['bucket_180_plus']}")
    print("Mapped column headers:")
    for field, info in report["column_headers"].items():
        print(f"  {field}: column {info['column']} = {info['header']!r}")
    print(f"Inconsistent cached age/bucket rows: {report['inconsistent_buckets']}")
    print(f"Extreme ages over 1000: {report['extreme_ages']}")
    for mapping in report["suffix_mappings"]:
        print(f"  row {mapping['source_row']}: {mapping['comment_id']}")
    for issue in report["inconsistent_rows"]:
        print(
            f"  row {issue['source_row']}: age={issue['age']}, "
            f"bucket={issue['aging_bucket']}, expected={issue['expected_bucket']}"
        )


def parse_args():
    parser = argparse.ArgumentParser(description="Import the real accrued-liabilities workbook into Smartsheet")
    parser.add_argument("--file", help="Path to the source Excel workbook")
    parser.add_argument("--sheet", default=DEFAULT_SHEET_NAME)
    parser.add_argument("--apply", action="store_true", help="Stage real rows after creating a backup")
    parser.add_argument("--replace", action="store_true", help="Required acknowledgement for staging replacement data")
    parser.add_argument("--reimport", action="store_true", help="Allow staging IDs that already exist (old rows are removed at finalize)")
    parser.add_argument("--finalize", metavar="STATE_FILE", help="Delete old rows and status column after staged verification")
    parser.add_argument("--confirm-replace", action="store_true")
    parser.add_argument("--rollback-stage", metavar="STATE_FILE")
    return parser.parse_args()


def main():
    args = parse_args()
    if args.finalize:
        if not args.confirm_replace:
            raise SystemExit("--finalize requires --confirm-replace")
        state = finalize_migration(args.finalize)
        print(f"Migration finalized. Backup: {state['backup_path']}")
        return
    if args.rollback_stage:
        state = rollback_staged_migration(args.rollback_stage)
        print(f"Staged rows removed. Existing rows were not changed. Backup: {state['backup_path']}")
        return
    if not args.file:
        raise SystemExit("--file is required for dry-run or staging")
    records, report = load_real_records(args.file, sheet_name=args.sheet)
    print_report(report)
    if not args.apply:
        print("Dry run only. No Smartsheet changes were made.")
        return
    if not args.replace:
        raise SystemExit("--apply requires --replace")
    state_path, state = stage_migration(args.file, sheet_name=args.sheet, allow_overlap=args.reimport)
    print(f"Staged {len(records)} rows without deleting existing rows.")
    print(f"Backup: {state['backup_path']}")
    print(f"State: {state_path}")
    print(f"Finalize only after review: python -m scripts.import_real_excel --finalize \"{state_path}\" --confirm-replace")


if __name__ == "__main__":
    main()
